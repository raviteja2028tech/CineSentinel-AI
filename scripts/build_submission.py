"""
Build the submission deliverables from the measured results.

Creates, under submission/:
  - CineSentinel_AI_Report.md     editable report source (generated)
  - CineSentinel_AI_Report.pdf    printed with headless Microsoft Edge
  - CineSentinel_AI_Presentation.pptx

Every figure is read from reports/metrics/ at build time, so the documents
cannot drift from the evaluated models. Team details are placeholders to be
filled in by hand.

Requires the extra packages in requirements-dev.txt (python-pptx, markdown)
and, for the PDF, Microsoft Edge.

Usage:  python scripts\\build_submission.py
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from cinesentinel.app_utils import (  # noqa: E402
    build_metrics_frame,
    load_dataset_stats,
    load_evaluation_results,
    load_feature_comparison,
    load_gb_diagnostics,
    load_gb_paper_result,
    load_training_summary,
)
from cinesentinel.config import (  # noqa: E402
    C_GRID,
    GB_PARAMS,
    MLP_PARAMS,
    MODEL_FAMILIES,
    RF_PARAMS,
    TFIDF_PARAMS,
)

SUBMISSION_DIR = PROJECT_ROOT / "submission"
SCREENSHOT_DIR = SUBMISSION_DIR / "screenshots"
REPORT_MD = SUBMISSION_DIR / "CineSentinel_AI_Report.md"
REPORT_PDF = SUBMISSION_DIR / "CineSentinel_AI_Report.pdf"
DECK_PPTX = SUBMISSION_DIR / "CineSentinel_AI_Presentation.pptx"

TITLE = "Machine Learning-Based Classification for Sentiment Analysis of IMDb Reviews"
APP = "CineSentinel AI — IMDb Sentiment Intelligence Lab"
COURSE = "UE24CS352A — Machine Learning"
TEAM_PLACEHOLDER = "[Team member 1 — SRN] · [Team member 2 — SRN] · [Team member 3 — SRN]"

REPRESENTATION_NAMES = {"binary": "binary", "count": "word-count", "tfidf": "TF-IDF"}

EDGE_CANDIDATES = [
    Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
]


# --------------------------------------------------------------------------
# Facts: everything the documents state, gathered once from the metrics files
# --------------------------------------------------------------------------


def gather_facts() -> dict:
    eval_results = load_evaluation_results()
    if not eval_results or len(eval_results) != 6:
        raise SystemExit("Need evaluation results for all six models - run cinesentinel.evaluate first.")
    summary = load_training_summary() or {}
    metrics = build_metrics_frame(eval_results, summary)
    dataset = load_dataset_stats()
    if dataset is None:
        raise SystemExit("Run scripts\\dataset_stats.py first.")

    gb_diag = load_gb_diagnostics()
    gb_paper = load_gb_paper_result()
    features = load_feature_comparison()

    facts = {
        "metrics": metrics,
        "dataset": dataset,
        "summary": summary,
        "best": metrics.iloc[0],
        "worst": metrics.iloc[-1],
        "top3": metrics.head(3),
        "top3_gap_pp": (metrics.iloc[0]["accuracy"] - metrics.iloc[2]["accuracy"]) * 100,
        "n_test": int(metrics.iloc[0][["tp", "fp", "fn", "tn"]].sum()),
        "gb_diag": gb_diag,
        "gb_paper": gb_paper,
        "features": features,
    }
    if gb_diag:
        stages = pd.DataFrame(gb_diag["stages"])
        at_100 = stages.loc[stages["n_estimators"] == 100, "validation_accuracy"]
        facts["gb_val_at_100"] = float(at_100.iloc[0]) if len(at_100) else None
    if features is not None:
        best_row = features.loc[features["validation_accuracy"].idxmax()]
        facts["feature_best"] = best_row
        # Is TF-IDF the best representation for every model and n-gram range?
        pivot = features.pivot_table(
            index=["model", "ngrams"], columns="representation", values="validation_accuracy"
        )
        facts["tfidf_always_best"] = bool((pivot.idxmax(axis=1) == "tfidf").all())
        facts["feature_range"] = (
            float(features["validation_accuracy"].min()),
            float(features["validation_accuracy"].max()),
        )
    return facts


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%"


def classifier_settings(summary: dict) -> list:
    """(model, family, settings) rows, using the hyperparameters actually selected."""
    lr_c = summary.get("logistic_regression", {}).get("best_C")
    svm_c = summary.get("linear_svm", {}).get("best_C")
    grid = "{" + ", ".join(f"{c:g}" for c in C_GRID) + "}"
    return [
        ("Logistic Regression", MODEL_FAMILIES["logistic_regression"], f"L2 penalty; C from {grid} by 3-fold CV → C = {lr_c:g}"),
        ("Linear SVM", MODEL_FAMILIES["linear_svm"], f"L2 penalty; same C grid → C = {svm_c:g}; Platt-calibrated"),
        ("Multinomial Naive Bayes", MODEL_FAMILIES["multinomial_nb"], "Laplace smoothing = 1"),
        (
            "Random Forest",
            MODEL_FAMILIES["random_forest"],
            f"{RF_PARAMS['n_estimators']} trees, Gini criterion, no depth limit",
        ),
        (
            "Gradient Boosting",
            MODEL_FAMILIES["gradient_boosting"],
            f"{GB_PARAMS['n_estimators']} trees (chosen on validation data), learning rate {GB_PARAMS['learning_rate']}",
        ),
        (
            "Neural network (MLPClassifier)",
            MODEL_FAMILIES["mlp_neural_net"],
            f"hidden layers {MLP_PARAMS['hidden_layer_sizes']}, logistic activation, Adam, early stopping",
        ),
    ]


# --------------------------------------------------------------------------
# Report (Markdown -> HTML -> PDF)
# --------------------------------------------------------------------------


def build_report_markdown(f: dict) -> str:
    d, m, best, worst = f["dataset"], f["metrics"], f["best"], f["worst"]
    overlap = d["cross_split_overlap"]
    top3_names = ", ".join(f["top3"]["model"].tolist()[:2]) + " and " + f["top3"]["model"].tolist()[2]

    results_rows = "\n".join(
        f"| {r['model']} | {r['accuracy']:.4f} | {r['positive_precision']:.4f} | {r['recall']:.4f} "
        f"| {r['f1']:.4f} | {r['macro_f1']:.4f} | {r['tp']:,} / {r['fp']:,} / {r['fn']:,} / {r['tn']:,} "
        f"| {r['train_seconds']:.1f} |"
        for _, r in m.iterrows()
    )
    settings_rows = "\n".join(f"| {name} | {settings} |" for name, _, settings in classifier_settings(f["summary"]))

    source_text = (
        "The files were downloaded from a HuggingFace mirror and all review–label pairs were verified "
        "to be identical to the official Stanford archive."
        if d["source"].get("verified_against_official")
        else "The files were obtained from the HuggingFace mirror of the official release."
    )

    gb_row = m.set_index("model_key").loc["gradient_boosting"]
    gb_text = ""
    if f["gb_diag"] and f["gb_paper"]:
        g = f["gb_diag"]
        gb_text = (
            f"With the paper's 100 trees, Gradient Boosting scored {pct(f['gb_paper']['test_metrics']['accuracy'])} "
            "on the test split. This was investigated rather than accepted: a control model trained on shuffled "
            f"labels scored {pct(g['shuffled_label_control_accuracy'], 1)} (chance), ruling out leakage or "
            f"misaligned labels, and even a {g['max_trees']}-tree model used only {g['features_used']:,} of "
            f"{g['n_features']:,} features, so the 100-tree model was simply too small. The tree count was therefore chosen on a validation split of the "
            f"training data ({g['fit_rows']:,} fit / {g['validation_rows']:,} validation), where accuracy rose "
            f"from {pct(f['gb_val_at_100'])} at 100 trees to {pct(g['selected_validation_accuracy'])} at "
            f"{g['selected_n_estimators']}. The final {g['selected_n_estimators']}-tree model, evaluated once, "
            f"scores {pct(gb_row['accuracy'])} on the test split."
        )

    feature_text = ""
    if f["features"] is not None:
        fb = f["feature_best"]
        lo, hi = f["feature_range"]
        feature_text = (
            "A side experiment compared binary, word-count and TF-IDF features with 1-, 1–2- and 1–3-grams for "
            "the three fast classifiers, on the same validation split. Validation accuracy ranged from "
            f"{pct(lo, 1)} to {pct(hi, 1)}; the best was {fb['model']} with {REPRESENTATION_NAMES[fb['representation']]} "
            f"{fb['ngrams']}s. "
            + (
                "TF-IDF gave the highest accuracy for every classifier and n-gram range, which differs from the "
                "paper, where binary features were best. "
                if f["tfidf_always_best"]
                else ""
            )
            + "Classifier settings were not re-tuned per representation."
        )

    return f"""# {TITLE}

**{APP}** · {COURSE}

**Team:** {TEAM_PLACEHOLDER}

## 1. Problem statement

Given the text of an IMDb movie review, classify the reviewer's opinion as **positive** or **negative**. This is document-level binary sentiment classification. The work follows the methodology of Wu and Shin [1], who compare six classifier families on this task.

## 2. Objectives

- Build a reproducible pipeline: data acquisition, text cleaning, feature extraction, training and evaluation.
- Train and compare six classifiers on identical features, using standard metrics on a held-out test set.
- Provide an interactive application for single-review prediction, model comparison and batch scoring of CSV files.

## 3. Dataset

The Stanford Large Movie Review Dataset [2] contains {d['total_rows']:,} labelled reviews in an official split of {d['train']['rows']:,} training and {d['test']['rows']:,} test reviews, each exactly half positive and half negative (rating ≥ 7 is positive, ≤ 4 is negative). The median review has {d['train']['words_median']:.0f} words. The official split is used unchanged and the test split is used only for the final evaluation. The public dataset contains {overlap['test_rows_also_in_train']} test reviews ({overlap['share_of_test'] * 100:.2f}%) that also occur in the training split; they were left in place, so test accuracy may be very slightly optimistic. {source_text}

## 4. Preprocessing and feature extraction

Each review is stripped of HTML, lowercased, stripped of punctuation and digits, tokenized, filtered for English stop words and lemmatized (NLTK). Reviews are then converted to **TF-IDF** vectors over unigrams and bigrams, limited to the {TFIDF_PARAMS['max_features']:,} most frequent terms (min. document frequency {TFIDF_PARAMS['min_df']}, sublinear term frequency). The vectorizer is fitted on the training split only, and the resulting matrix is shared by all six models.

## 5. Classifiers

| Model | Settings |
|---|---|
{settings_rows}

Settings follow §3.3 of [1] except the Gradient Boosting tree count (Section 7). All model selection used training data only: cross-validation for C, and a held-out 10% of the training split for the neural network's early stopping.

## 6. Implementation

Python 3.11 with scikit-learn [3], NLTK, Streamlit and Plotly, running on a laptop CPU with no GPU or paid service. Training and evaluation are command-line modules that write model files and metrics files; the five-page Streamlit dashboard reads those files, so no figure is typed into the application. The project has an automated test suite (pytest).

## 7. Results

Measured once on the {f['n_test']:,}-review test split. The positive class is "positive"; precision, recall and F1 are for that class, and Macro-F1 averages both classes.

| Model | Accuracy | Precision | Recall | F1 | Macro-F1 | TP / FP / FN / TN | Train (s) |
|---|---|---|---|---|---|---|---|
{results_rows}

{best['model']} has the highest measured accuracy ({pct(best['accuracy'])}), but the top three ({top3_names}) are within {f['top3_gap_pp']:.2f} percentage points of one another. Each model was trained once, so this gap is too small to rank them reliably. {worst['model']} is the weakest model ({pct(worst['accuracy'])}).

{gb_text}

{feature_text}

## 8. Conclusion

All six classifiers were implemented, trained and evaluated on the official test split. The two linear models and the neural network reach about {best['accuracy'] * 100:.0f}% accuracy, while Naive Bayes and the tree ensembles are lower; this is the same overall pattern reported in [1]. For this task a linear model offers the best balance: it matches the neural network's accuracy with a small fraction of the training time and remains interpretable.

## 9. Limitations

- Not an exact reproduction of [1], which uses 3-gram features and does not publish its vocabulary size, preprocessing code or random seeds.
- Each model was trained once; no confidence intervals or repeated seeds are reported.
- Bag-of-words features ignore word order, so sarcasm and long-range negation are handled poorly.
- The feature-representation experiment covers three classifiers on validation data, not all six on the test split.

## References

[1] C.-L. Wu and S.-L. Shin, "Machine Learning based classification for Sentimental analysis of IMDb reviews," Stanford University, CS229 project report.

[2] A. L. Maas, R. E. Daly, P. T. Pham, D. Huang, A. Y. Ng and C. Potts, "Learning Word Vectors for Sentiment Analysis," *Proc. 49th Annual Meeting of the ACL: Human Language Technologies*, 2011. Dataset: https://ai.stanford.edu/~amaas/data/sentiment/

[3] F. Pedregosa et al., "Scikit-learn: Machine Learning in Python," *Journal of Machine Learning Research*, vol. 12, pp. 2825–2830, 2011.
"""


REPORT_CSS = """
@page { size: A4; margin: 13mm 14mm 13mm 14mm; }
* { box-sizing: border-box; }
body { font-family: 'Segoe UI', Calibri, Arial, sans-serif; font-size: 9.1pt; line-height: 1.36; color: #111827; margin: 0; }
h1 { font-size: 15.5pt; line-height: 1.2; margin: 0 0 3pt 0; color: #1E1B4B; }
h1 + p { margin: 0 0 2pt 0; color: #374151; }
h2 { font-size: 10.6pt; margin: 8pt 0 2.5pt 0; padding-bottom: 1.5pt; color: #312E81; border-bottom: 0.8pt solid #C7D2FE; break-after: avoid; }
p { margin: 0 0 4pt 0; text-align: justify; }
ul { margin: 0 0 4pt 0; padding-left: 14pt; }
li { margin: 0 0 1pt 0; }
table { border-collapse: collapse; width: 100%; margin: 3pt 0 5pt 0; font-size: 8.4pt; break-inside: avoid; }
th { background: #EEF2FF; color: #1E1B4B; text-align: left; padding: 2.4pt 4pt; border: 0.6pt solid #C7D2FE; }
td { padding: 2.2pt 4pt; border: 0.6pt solid #D1D5DB; vertical-align: top; }
td:not(:first-child), th:not(:first-child) { white-space: nowrap; }
table.wrap td:not(:first-child) { white-space: normal; }
a { color: #3730A3; text-decoration: none; word-break: break-all; }
"""


def build_report(f: dict) -> None:
    import markdown

    text = build_report_markdown(f)
    REPORT_MD.write_text(text, encoding="utf-8")

    body = markdown.markdown(text, extensions=["tables"])
    # The settings table holds long text in its second column and must wrap.
    body = body.replace("<table>", "<table class='wrap'>", 1)
    html = (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        f"<title>{TITLE}</title><style>{REPORT_CSS}</style></head><body>{body}</body></html>"
    )
    edge = next((p for p in EDGE_CANDIDATES if p.exists()), None)
    if edge is None:
        print("Microsoft Edge not found: wrote the .md only; the PDF was not rebuilt.")
        return
    if REPORT_PDF.exists():
        REPORT_PDF.unlink()
    # The HTML and the browser profile are scratch files, kept out of the project.
    work = Path(tempfile.mkdtemp(prefix="cinesentinel-report-"))
    try:
        html_path = work / "report.html"
        html_path.write_text(html, encoding="utf-8")
        subprocess.run(
            [
                str(edge), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                f"--user-data-dir={work / 'profile'}", f"--print-to-pdf={REPORT_PDF}", html_path.as_uri(),
            ],
            check=False, capture_output=True, timeout=120,
        )
    finally:
        shutil.rmtree(work, ignore_errors=True)
    print(f"report: {REPORT_PDF} ({'ok' if REPORT_PDF.exists() else 'NOT CREATED'})")


# --------------------------------------------------------------------------
# Presentation
# --------------------------------------------------------------------------


def build_deck(f: dict) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    NAVY, VIOLET, CYAN = RGBColor(0x1E, 0x1B, 0x4B), RGBColor(0x6D, 0x4A, 0xF0), RGBColor(0x08, 0x91, 0xB2)
    INK, MUTED, LIGHT = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x64, 0x74, 0x8B), RGBColor(0xEE, 0xF2, 0xFF)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    FONT = "Segoe UI"

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    def text_box(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = Inches(0.05)
        frame.margin_top = frame.margin_bottom = Inches(0.03)
        return frame

    def write(frame, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, first=False, space_after=4):
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
        run.font.color.rgb = color
        return paragraph

    def bullets(frame, items, size=16, color=INK, first=True, space_after=7, marker="•"):
        for i, item in enumerate(items):
            label = f"{i + 1}." if marker == "number" else marker
            paragraph = write(
                frame, f"{label}\t{item}", size=size, color=color,
                first=(first and i == 0), space_after=space_after,
            )
            # Hanging indent: wrapped lines line up with the text, not the marker.
            properties = paragraph._p.get_or_add_pPr()
            indent = int(Pt(size) * 1.25)
            properties.set("marL", str(indent))
            properties.set("indent", str(-indent))

    def rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, left, top, width, height)
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(1)
        s.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.09
        return s

    def new_slide(title, number):
        slide = prs.slides.add_slide(blank)
        rect(slide, 0, 0, W, Inches(0.12), VIOLET)
        write(text_box(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8)), title, size=30, bold=True, color=NAVY, first=True)
        rect(slide, Inches(0.65), Inches(1.15), Inches(1.1), Inches(0.05), CYAN)
        footer = text_box(slide, Inches(0.6), Inches(7.02), Inches(12.1), Inches(0.35))
        write(footer, f"CineSentinel AI  ·  {COURSE}", size=10, color=MUTED, first=True)
        write(text_box(slide, Inches(12.0), Inches(7.02), Inches(0.8), Inches(0.35)), str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT, first=True)
        return slide

    def heading(slide, left, top, width, text):
        write(text_box(slide, left, top, width, Inches(0.45)), text, size=17, bold=True, color=VIOLET, first=True)

    def table(slide, left, top, width, rows, col_widths, font_size=12, row_height=0.4):
        shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, Inches(row_height * len(rows)))
        tbl = shape.table
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(width * w))
        for r, row in enumerate(rows):
            tbl.rows[r].height = Inches(row_height)
            for c, value in enumerate(row):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = str(value)
                run.font.size, run.font.name, run.font.bold = Pt(font_size), FONT, (r == 0)
                run.font.color.rgb = WHITE if r == 0 else INK
        return tbl

    d, m, best, worst = f["dataset"], f["metrics"], f["best"], f["worst"]
    overlap = d["cross_split_overlap"]

    # ---- 1. Title
    slide = prs.slides.add_slide(blank)
    rect(slide, 0, 0, W, prs.slide_height, NAVY)
    rect(slide, 0, Inches(5.05), W, Inches(0.06), CYAN)
    frame = text_box(slide, Inches(0.9), Inches(1.2), Inches(11.5), Inches(3.6), anchor=MSO_ANCHOR.MIDDLE)
    write(frame, "CINESENTINEL AI  ·  IMDb SENTIMENT INTELLIGENCE LAB", size=14, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC), first=True, space_after=14)
    write(frame, TITLE, size=38, bold=True, color=WHITE, space_after=14)
    write(frame, COURSE + "  ·  Mini-project", size=18, color=RGBColor(0xC7, 0xD2, 0xFE))
    frame = text_box(slide, Inches(0.9), Inches(5.35), Inches(11.5), Inches(1.6))
    write(frame, "Team", size=13, bold=True, color=RGBColor(0x67, 0xE8, 0xF9), first=True)
    write(frame, "[Team member 1 — SRN]     [Team member 2 — SRN]     [Team member 3 — SRN]", size=16, color=WHITE)
    write(frame, "[Section / Guide / Date]", size=13, color=RGBColor(0xC7, 0xD2, 0xFE))

    # ---- 2. Problem and objectives
    slide = new_slide("Problem statement and objectives", 2)
    heading(slide, Inches(0.6), Inches(1.5), Inches(5.8), "Problem")
    frame = text_box(slide, Inches(0.6), Inches(2.0), Inches(5.8), Inches(4.6))
    bullets(frame, [
        "Decide whether an IMDb movie review is positive or negative",
        "Document-level binary text classification",
        "Reference: Wu & Shin (Stanford CS229) compare six classifier families on this task",
        "Question: how do the six compare in accuracy and in training cost?",
    ], size=19, space_after=12)
    heading(slide, Inches(6.9), Inches(1.5), Inches(5.8), "Objectives")
    frame = text_box(slide, Inches(6.9), Inches(2.0), Inches(5.8), Inches(4.6))
    bullets(frame, [
        "Reproducible pipeline: data → cleaning → features → training → evaluation",
        "Train six classifiers on identical features",
        "Evaluate with standard metrics on the official held-out test split",
        "Interactive app: live prediction, model comparison, batch CSV scoring",
    ], size=19, space_after=12)

    # ---- 3. Dataset and preprocessing
    slide = new_slide("Dataset and preprocessing", 3)
    kpis = [
        (f"{d['total_rows']:,}", "labelled reviews"),
        (f"{d['train']['rows'] // 1000}k / {d['test']['rows'] // 1000}k", "official train / test split"),
        ("50% / 50%", "positive / negative"),
        (f"{d['train']['words_median']:.0f}", "median words per review"),
    ]
    for i, (value, label) in enumerate(kpis):
        left = Inches(0.6 + i * 3.08)
        rect(slide, left, Inches(1.5), Inches(2.9), Inches(1.25), LIGHT, line=RGBColor(0xC7, 0xD2, 0xFE), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        frame = text_box(slide, left, Inches(1.56), Inches(2.9), Inches(1.15), anchor=MSO_ANCHOR.MIDDLE)
        write(frame, value, size=26, bold=True, color=NAVY, align=PP_ALIGN.CENTER, first=True, space_after=0)
        write(frame, label, size=12, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)
    heading(slide, Inches(0.6), Inches(3.0), Inches(5.8), "Stanford Large Movie Review Dataset")
    frame = text_box(slide, Inches(0.6), Inches(3.5), Inches(5.8), Inches(3.4))
    bullets(frame, [
        "Positive = rating ≥ 7, negative = rating ≤ 4",
        "Official split used unchanged; test split used only for final evaluation",
        f"{overlap['test_rows_also_in_train']} test reviews ({overlap['share_of_test'] * 100:.2f}%) also occur in train — a property of the public data, left in place and disclosed",
        *(["Content verified identical to the official Stanford archive"] if d["source"].get("verified_against_official") else []),
    ], size=15)
    heading(slide, Inches(6.9), Inches(3.0), Inches(5.8), "Preprocessing and features")
    frame = text_box(slide, Inches(6.9), Inches(3.5), Inches(5.8), Inches(3.4))
    bullets(frame, [
        "Strip HTML → lowercase → remove punctuation and digits",
        "Remove stop words → lemmatize (NLTK)",
        f"TF-IDF, unigrams + bigrams, {TFIDF_PARAMS['max_features']:,} terms",
        "Vectorizer fitted on the training split only",
    ], size=15)

    # ---- 4. Methodology and architecture
    slide = new_slide("Methodology and system architecture", 4)

    def flow(labels, top, fill, line, text_color=INK):
        n = len(labels)
        box_w, gap = Inches(2.15), Inches(0.36)
        start = int((W - (box_w * n + gap * (n - 1))) / 2)
        boxes = []
        for i, (title, sub) in enumerate(labels):
            left = start + i * (box_w + gap)
            rect(slide, left, top, box_w, Inches(1.05), fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            frame = text_box(slide, left, top, box_w, Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
            write(frame, title, size=14, bold=True, color=text_color, align=PP_ALIGN.CENTER, first=True, space_after=1)
            write(frame, sub, size=11, color=text_color if text_color == WHITE else MUTED, align=PP_ALIGN.CENTER, space_after=0)
            boxes.append((left, box_w))
        for (left, width), (next_left, _) in zip(boxes, boxes[1:]):
            arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left + width, top + Inches(0.525), next_left, top + Inches(0.525))
            arrow.line.color.rgb = MUTED
            arrow.line.width = Pt(1.75)
            # python-pptx has no arrowhead API; add the tail-end arrow in XML.
            ln = arrow.line._get_or_add_ln()
            ln.append(ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"}))

    heading(slide, Inches(0.6), Inches(1.45), Inches(8), "Offline: training and evaluation")
    flow(
        [
            ("Dataset", "25k train / 25k test"),
            ("Cleaning", "preprocessing.py"),
            ("TF-IDF", "features.py"),
            ("Training", "train.py · six models"),
            ("Evaluation", "evaluate.py · test split"),
        ],
        Inches(1.95), LIGHT, RGBColor(0xC7, 0xD2, 0xFE),
    )
    heading(slide, Inches(0.6), Inches(3.35), Inches(8), "Online: Streamlit dashboard")
    flow(
        [
            ("Saved files", "models + metrics"),
            ("Inference", "predict.py"),
            ("Dashboard", "5 pages · Plotly charts"),
        ],
        Inches(3.85), NAVY, NAVY, text_color=WHITE,
    )
    frame = text_box(slide, Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.7))
    bullets(frame, [
        "All six models share one feature matrix, so the comparison is like-for-like",
        "Model selection uses training data only (cross-validation, validation split); the test split is evaluated once",
        "The dashboard reads saved metrics files — no figure is typed into the app. Runs on a laptop CPU, no GPU or paid API",
    ], size=14, space_after=5)

    # ---- 5. Six classifiers
    slide = new_slide("Six machine-learning classifiers", 5)
    rows = [("Model", "Family", "Key settings")] + [tuple(r) for r in classifier_settings(f["summary"])]
    table(slide, Inches(0.6), Inches(1.5), Inches(12.1), rows, [0.24, 0.15, 0.61], font_size=13, row_height=0.58)
    frame = text_box(slide, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.2))
    bullets(frame, [
        "Settings follow §3.3 of the reference paper, except the Gradient Boosting tree count",
        "The SVM is Platt-calibrated so that it can report a probability rather than a raw decision score",
    ], size=14, space_after=4)

    # ---- 6. Results
    slide = new_slide("Experimental results", 6)
    chart_data = CategoryChartData()
    chart_data.categories = m["short_name"].tolist()
    chart_data.add_series("Accuracy", [round(v * 100, 2) for v in m["accuracy"]])
    chart_data.add_series("F1 (positive class)", [round(v * 100, 2) for v in m["f1"]])
    frame_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.4), Inches(6.2), Inches(4.5), chart_data)
    chart = frame_shape.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size, chart.legend.font.name = Pt(11), FONT
    axis = chart.value_axis
    axis.minimum_scale = float(int(min(m["accuracy"].min(), m["f1"].min()) * 100 / 5) * 5 - 5)
    axis.maximum_scale = 95.0
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    axis.tick_labels.font.size = Pt(11)
    axis.tick_labels.number_format = '0"%"'
    axis.tick_labels.number_format_is_linked = False
    chart.category_axis.tick_labels.font.size = Pt(10)
    for series, color in zip(chart.plots[0].series, (VIOLET, CYAN)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.overlap = -10

    rows = [("Model", "Acc.", "Prec.", "Recall", "F1", "Train")]
    for _, r in m.iterrows():
        rows.append((
            r["short_name"], f"{r['accuracy'] * 100:.2f}", f"{r['positive_precision'] * 100:.2f}",
            f"{r['recall'] * 100:.2f}", f"{r['f1'] * 100:.2f}", f"{r['train_seconds']:.0f} s" if r["train_seconds"] >= 10 else f"{r['train_seconds']:.1f} s",
        ))
    table(slide, Inches(6.95), Inches(1.5), Inches(5.85), rows, [0.27, 0.14, 0.14, 0.15, 0.14, 0.16], font_size=12, row_height=0.42)
    write(text_box(slide, Inches(6.95), Inches(4.5), Inches(5.85), Inches(0.35)),
          f"Percent, on {f['n_test']:,} test reviews. Chart axis does not start at zero.", size=10, color=MUTED, first=True)
    frame = text_box(slide, Inches(6.95), Inches(4.85), Inches(5.85), Inches(2.1))
    points = [
        f"Highest: {best['model']} ({pct(best['accuracy'])})",
        f"Top three within {f['top3_gap_pp']:.2f} points — one run each, so effectively tied",
        f"Weakest: {worst['model']} ({pct(worst['accuracy'])})",
    ]
    bullets(frame, points, size=13, space_after=4)

    # ---- 7. Application
    slide = new_slide("Application interface and demo workflow", 7)
    shots = [("live.png", "Live Sentiment Analyzer"), ("comparison.png", "Model Comparison Lab")]
    available = [(SCREENSHOT_DIR / name, caption) for name, caption in shots if (SCREENSHOT_DIR / name).exists()]
    pic_w, pic_h = Inches(4.35), Inches(2.72)
    for i, (path, caption) in enumerate(available):
        left = Inches(0.6 + i * 4.55)
        picture = slide.shapes.add_picture(str(path), left, Inches(1.5), width=pic_w, height=pic_h)
        picture.line.color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
        picture.line.width = Pt(1)
        write(text_box(slide, left, Inches(4.26), pic_w, Inches(0.35)), caption, size=11, color=MUTED, align=PP_ALIGN.CENTER, first=True)
    if not available:
        frame = text_box(slide, Inches(0.6), Inches(1.5), Inches(8.9), Inches(2.7), anchor=MSO_ANCHOR.MIDDLE)
        write(frame, "[Insert application screenshots here]", size=18, color=MUTED, align=PP_ALIGN.CENTER, first=True)
    heading(slide, Inches(9.75), Inches(1.4), Inches(3.1), "Five pages")
    frame = text_box(slide, Inches(9.75), Inches(1.9), Inches(3.1), Inches(2.7))
    bullets(frame, ["Home / overview", "Live Sentiment Analyzer", "Model Comparison Lab", "Batch Prediction Studio", "Methodology"], size=13, space_after=3)
    heading(slide, Inches(0.6), Inches(4.75), Inches(8), "Demo workflow")
    frame = text_box(slide, Inches(0.6), Inches(5.22), Inches(12.1), Inches(1.75))
    bullets(frame, [
        "Home: dataset statistics and the measured leaderboard",
        "Live analyzer: paste a review, pick a model, compare all six verdicts and see any disagreement",
        "Comparison lab: metrics, confusion matrices, ROC curves, training time",
        "Batch studio: upload a CSV, score it, download the predictions",
    ], size=13, space_after=3, marker="number")

    # ---- 8. Conclusion
    slide = new_slide("Conclusion, limitations and future scope", 8)
    columns = [
        ("Conclusion", [
            "All six classifiers implemented, trained and evaluated on the official test split",
            f"Linear models and the neural network reach about {best['accuracy'] * 100:.0f}% accuracy",
            "A linear model gives the best balance of accuracy, training time and interpretability",
        ]),
        ("Limitations", [
            "Not an exact reproduction of the paper (1–2-grams, not 3-grams)",
            "One training run per model — no confidence intervals",
            "Bag-of-words features ignore word order",
        ]),
        ("Future scope", [
            "Repeat runs over several seeds",
            "Tune the tree ensembles further",
            "Binary and word-count features for all six models",
            "Sequence models: RNN, transformer",
        ]),
    ]
    for i, (title, items) in enumerate(columns):
        left = Inches(0.6 + i * 4.1)
        rect(slide, left, Inches(1.5), Inches(3.9), Inches(4.3), LIGHT, line=RGBColor(0xC7, 0xD2, 0xFE), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        heading(slide, left + Inches(0.2), Inches(1.65), Inches(3.5), title)
        frame = text_box(slide, left + Inches(0.2), Inches(2.2), Inches(3.5), Inches(3.5))
        bullets(frame, items, size=15, space_after=10)

    # python-pptx stamps its template author on new files; replace it.
    properties = prs.core_properties
    properties.title = TITLE
    properties.subject = COURSE
    properties.author = "CineSentinel AI project team"
    properties.last_modified_by = "CineSentinel AI project team"
    prs.save(DECK_PPTX)
    print(f"deck: {DECK_PPTX} ({len(prs.slides)} slides)")


def main() -> None:
    SUBMISSION_DIR.mkdir(exist_ok=True)
    facts = gather_facts()
    build_report(facts)
    build_deck(facts)


if __name__ == "__main__":
    main()
